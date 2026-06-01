"""
Kautilya AI — Export Routes Blueprint
Handles /api/export/* endpoints (PDF, DOCX, Excel).
"""
import io
import re
from flask import Blueprint, request, jsonify, send_file

export_bp = Blueprint('export', __name__)


@export_bp.route('/export/pdf', methods=['POST'])
def export_pdf():
    try:
        data = request.json
        chat_history = data.get('history', [])
        
        try:
            from fpdf import FPDF
        except ImportError:
            return jsonify({"error": "fpdf library not installed. Please install fpdf2."}), 500
            
        class PDF(FPDF):
            def header(self): pass
            def footer(self): pass

        pdf = PDF()
        pdf.add_page()
        pdf.set_font("Arial", size=11)
        
        target_msg = None
        for msg in reversed(chat_history):
            if msg.get('role', 'Unknown').lower() == 'assistant':
                target_msg = msg
                break
                
        if target_msg:
            content = target_msg.get('content', '')
            if isinstance(content, list):
                content = " ".join([p.get('text', '') for p in content if p.get('type') == 'text'])
            content = re.sub(r'\[EXPORT:.*?\]', '', content, flags=re.IGNORECASE).strip()
            
            if content:
                safe_content = content.encode('latin-1', 'replace').decode('latin-1')
                pdf.multi_cell(0, 6, safe_content)
                pdf.ln(5)

        buffer = io.BytesIO()
        # fpdf2: output() returns a bytearray; older fpdf returned a latin-1 str.
        out = pdf.output()
        if isinstance(out, str):
            out = out.encode('latin-1')
        buffer.write(bytes(out))
        buffer.seek(0)
        
        return send_file(buffer, as_attachment=True, download_name='content_export.pdf', mimetype='application/pdf')
    except Exception as e:
        print(f"PDF Export failed: {e}")
        return jsonify({"error": str(e)}), 500


@export_bp.route('/export/docx', methods=['POST'])
def export_docx():
    try:
        data = request.json
        chat_history = data.get('history', [])
        
        try:
            from docx import Document
            from docx.shared import Pt
        except ImportError:
            return jsonify({"error": "python-docx library not installed."}), 500
            
        document = Document()
        
        target_msg = None
        for msg in reversed(chat_history):
            if msg.get('role', 'Unknown').lower() == 'assistant':
                target_msg = msg
                break
                
        if target_msg:
            content = target_msg.get('content', '')
            if isinstance(content, list):
                content = " ".join([p.get('text', '') for p in content if p.get('type') == 'text'])
            content = re.sub(r'\[EXPORT:.*?\]', '', content, flags=re.IGNORECASE).strip()
            
            if content:
                p = document.add_paragraph()
                run = p.add_run(content)
                run.font.size = Pt(11)

        buffer = io.BytesIO()
        document.save(buffer)
        buffer.seek(0)
        
        return send_file(buffer, as_attachment=True, download_name='content_export.docx', 
                         mimetype='application/vnd.openxmlformats-officedocument.wordprocessingml.document')
    except Exception as e:
         print(f"DOCX Export failed: {e}")
         return jsonify({"error": str(e)}), 500


@export_bp.route('/export/excel', methods=['POST'])
def export_excel():
    try:
        data = request.json
        chat_history = data.get('history', [])
        
        try:
            import openpyxl
        except ImportError:
            return jsonify({"error": "openpyxl library not installed."}), 500
            
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.append(["Content"])
        
        target_msg = None
        for msg in reversed(chat_history):
            if msg.get('role', 'Unknown').lower() == 'assistant':
                target_msg = msg
                break
                
        if target_msg:
            content = target_msg.get('content', '')
            if isinstance(content, list):
               content = " ".join([p.get('text', '') for p in content if p.get('type') == 'text'])
            content = re.sub(r'\[EXPORT:.*?\]', '', content, flags=re.IGNORECASE).strip()
            if content:
                ws.append([content])

        buffer = io.BytesIO()
        wb.save(buffer)
        buffer.seek(0)
        
        return send_file(buffer, as_attachment=True, download_name='content_export.xlsx', 
                         mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    except Exception as e:
         print(f"Excel Export failed: {e}")
         return jsonify({"error": str(e)}), 500
